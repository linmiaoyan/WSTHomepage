"""
文件处理服务
"""
import os
from pathlib import Path
from typing import List, Dict, Any
from docx import Document
from openpyxl import load_workbook
import pandas as pd
import sys
from pathlib import Path
sys.path.insert(0, str(Path(__file__).parent.parent.parent))
import config


def extract_placeholders_from_docx(file_path: str) -> List[str]:
    """
    从Word文档中提取占位符
    占位符格式：{{field_name}}
    """
    placeholders = []
    doc = Document(file_path)
    
    # 提取段落中的占位符
    for paragraph in doc.paragraphs:
        import re
        matches = re.findall(r'\{\{(\w+)\}\}', paragraph.text)
        placeholders.extend(matches)
    
    # 提取表格中的占位符
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                matches = re.findall(r'\{\{(\w+)\}\}', cell.text)
                placeholders.extend(matches)
    
    return list(set(placeholders))  # 去重


def extract_placeholders_from_xlsx(file_path: str) -> List[str]:
    """
    从Excel文件中提取占位符
    """
    placeholders = []
    wb = load_workbook(file_path, data_only=False)
    
    import re
    for sheet_name in wb.sheetnames:
        sheet = wb[sheet_name]
        for row in sheet.iter_rows():
            for cell in row:
                if cell.value and isinstance(cell.value, str):
                    matches = re.findall(r'\{\{(\w+)\}\}', cell.value)
                    placeholders.extend(matches)
    
    return list(set(placeholders))


def extract_placeholders(file_path: str) -> List[str]:
    """根据文件类型提取占位符"""
    ext = Path(file_path).suffix.lower()
    if ext in ['.docx', '.doc']:
        return extract_placeholders_from_docx(file_path)
    elif ext in ['.xlsx', '.xls']:
        return extract_placeholders_from_xlsx(file_path)
    else:
        return []


# ========== 大模型集成部分（可选） ==========
def llm_extract_table_structure(file_path: str) -> Dict[str, Any]:
    """
    使用大模型智能提取表格结构
    可以自动识别哪些位置需要填充，无需手动标注占位符
    
    需要多模态大模型（如GPT-4V、Claude Vision）
    """
    # TODO: 集成大模型
    # 1. 将文件转换为图片
    # 2. 调用视觉-语言模型分析图片
    # 3. 返回字段映射关系
    # 
    # 示例：
    # import base64
    # from PIL import Image
    # 
    # # 将文档转换为图片
    # image = convert_doc_to_image(file_path)
    # image_base64 = base64.b64encode(image).decode()
    # 
    # # 调用大模型
    # response = llm_api.analyze_document(
    #     image=image_base64,
    #     prompt="识别这个表格中需要填充的字段位置"
    # )
    # 
    # return {
    #     "fields": ["name", "sex", "phone"],
    #     "positions": [...]
    # }
    return {}


def llm_auto_annotate_template(file_path: str) -> str:
    """
    使用大模型自动标注模板文件
    自动识别需要填充的位置并添加占位符
    """
    # TODO: 集成大模型
    # 1. 分析文档结构
    # 2. 识别需要填充的字段
    # 3. 自动添加占位符（如{{name}}）
    # 4. 保存标注后的文件
    return file_path

